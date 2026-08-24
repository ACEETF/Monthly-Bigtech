# -*- coding: utf-8 -*-
"""iM증권 ACE ETF 프로모션 참가 시행문(기안지) - ACE 고배당주Plus커버드콜액티브 외 4종
근거 공문: 한국투자신탁운용-2026-2275 (2026-06-16, iM증권 고객대상 ACE ETF 프로모션 참가)
예상비용: ISA적립식 200만 + 연금저축 200만 + ISA 순매수 449만 + DC/IRP 순매수 449만 = 12,980,000원
(아메리카노 단가 4,900원 기준, 한국투자신탁운용 전액 부담 → 총괄본부장 전결)
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "맑은 고딕"
L, C, R = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT


def sf(run, size=10, bold=False, color=None):
    run.font.size = Pt(size); run.bold = bold
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), FONT)
    if color: run.font.color.rgb = color


def para(doc, text="", size=10, bold=False, align=None, after=3, color=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.2
    if text: sf(p.add_run(text), size, bold, color)
    return p


def bd(cell):
    tcPr = cell._tc.get_or_add_tcPr(); tcB = tcPr.find(qn('w:tcBorders'))
    if tcB is None: tcB = OxmlElement('w:tcBorders'); tcPr.append(tcB)
    for e in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000'); tcB.append(el)


def ct(cell, text, size=9, bold=False, align=C, shade=None):
    cell.text = ""
    for i, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        sf(p.add_run(line), size, bold)
    cell.vertical_alignment = 1
    if shade:
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), shade); cell._tc.get_or_add_tcPr().append(sh)
    bd(cell)


def ab(t):
    for row in t.rows:
        for c in row.cells: bd(c)


def fixw(t, widths):
    t.autofit = False; t.allow_autofit = False
    tblPr = t._tbl.tblPr; lay = tblPr.find(qn('w:tblLayout'))
    if lay is None: lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
    lay.set(qn('w:type'), 'fixed')
    for row in t.rows:
        for c, w in zip(row.cells, widths): c.width = w
    for col, w in zip(t.columns, widths): col.width = w


doc = Document(); s = doc.sections[0]
s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.5); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

para(doc, "기 안 지", size=16, bold=True, align=C, after=10)
h = doc.add_table(rows=4, cols=6); h.alignment = WD_TABLE_ALIGNMENT.CENTER
ct(h.rows[0].cells[0], "문서번호", bold=True, shade="F2F2F2"); ct(h.rows[0].cells[1], "ETF마케팅부-2026-____", align=L)
h.rows[0].cells[1].merge(h.rows[0].cells[2])
ct(h.rows[0].cells[3], "기안일자", bold=True, shade="F2F2F2"); ct(h.rows[0].cells[4], "2026-__-__", align=L)
h.rows[0].cells[4].merge(h.rows[0].cells[5])
ct(h.rows[1].cells[0], "기안부서", bold=True, shade="F2F2F2"); ct(h.rows[1].cells[1], "ETF마케팅부", align=L)
ct(h.rows[1].cells[2], "기 안 자", bold=True, shade="F2F2F2"); ct(h.rows[1].cells[3], "채보라[부서원]", align=L)
ct(h.rows[1].cells[4], "전화번호", bold=True, shade="F2F2F2"); ct(h.rows[1].cells[5], "02-2055-5348", align=L)
ct(h.rows[2].cells[0], "전결근거", bold=True, shade="F2F2F2")
ct(h.rows[2].cells[1], "1. 관리공통업무  5. 경비예산집행  라. 광고선전비, 판매부대비, 행사비  3,000만원이하 (총괄본부장 전결)", align=L)
h.rows[2].cells[1].merge(h.rows[2].cells[5])
ct(h.rows[3].cells[0], "제    목", bold=True, shade="F2F2F2")
ct(h.rows[3].cells[1], "iM증권 고객대상 ACE ETF 프로모션 참가 시행의 건 (ACE 고배당주Plus커버드콜액티브 외 4종목)", align=L)
h.rows[3].cells[1].merge(h.rows[3].cells[5])
ab(h); fixw(h, [Cm(2.2), Cm(3.0), Cm(2.2), Cm(3.2), Cm(2.2), Cm(4.6)])
para(doc, "", after=8)

para(doc, "        한국투자신탁운용-2026-2275 <iM증권 고객대상 ACE ETF 프로모션 참가> 공문에 따라, "
          "아래와 같이 프로모션 시행을 요청하오니 재가하여 주시기 바랍니다.", align=L, after=8)
para(doc, "- 아    래 -", bold=True, align=C, after=8)
para(doc, "1. 목적 : ACE ETF 상품에 대한 투자자의 인지도 제고 및 거래 활성화", after=6)
para(doc, "2. 주요내용", after=4)

m = doc.add_table(rows=4, cols=2); m.alignment = WD_TABLE_ALIGNMENT.CENTER
ct(m.rows[0].cells[0], "구 분", size=10, bold=True, shade="F2F2F2"); ct(m.rows[0].cells[1], "세부 내용", size=10, bold=True, shade="F2F2F2")
ct(m.rows[1].cells[0], "기 간", size=10, bold=True); ct(m.rows[1].cells[1], "2026년 7월 1일 ~ 2026년 9월 30일", size=10, align=L)
ct(m.rows[2].cells[0], "프로모션\n내용", size=10, bold=True)
ct(m.rows[3].cells[0], "예상\n비용", size=10, bold=True)

box = m.rows[2].cells[1]; box.text = ""


def bl(text, size=10, bold=False, after=2, first=False):
    p = box.paragraphs[0] if first else box.add_paragraph()
    p.alignment = L; p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.15
    sf(p.add_run(text), size, bold); return p


bl("1. 대상고객 : iM증권 연금저축/ISA/DC/IRP 계좌 가입고객", first=True)
bl("2. 대상상품 : ACE 고배당주Plus커버드콜액티브, ACE K휴머노이드로봇산업Top2+,", after=0)
bl("                     ACE 원자력Top10, ACE K반도체Top2+, ACE 코리아AI전력Top10", after=4)
bl("3. 주요내용 :", after=3)
ev = box.add_table(rows=5, cols=3)
for j, hh in enumerate(["구분", "주요내용", "기타"]): ct(ev.rows[0].cells[j], hh, bold=True, shade="F2F2F2")
ct(ev.rows[1].cells[0], "ISA 적립식", bold=True)
ct(ev.rows[1].cells[1], "- 대상상품 ETF 적립식 등록 (1년, 1회당 10만원 이상)\n- 1만원 모바일 문화상품권 증정 (200명 한정)", align=L)
ct(ev.rows[2].cells[0], "연금저축\n적립식", bold=True)
ct(ev.rows[2].cells[1], "- 대상상품 ETF 적립식 등록 (1년, 1회당 10만원 이상)\n- 1만원 모바일 문화상품권 증정 (200명 한정)", align=L)
ct(ev.rows[3].cells[0], "ISA 순매수", bold=True)
ct(ev.rows[3].cells[1], "- 대상 ETF 순매수 금액에 따라 경품 증정 (300명 한정)", align=L)
ct(ev.rows[4].cells[0], "DC/IRP\n순매수", bold=True)
ct(ev.rows[4].cells[1], "- 대상 ETF 순매수 금액에 따라 경품 증정 (300명 한정)", align=L)
ct(ev.rows[1].cells[2], "- 운용사 비용 부담")
ev.rows[1].cells[2].merge(ev.rows[2].cells[2]).merge(ev.rows[3].cells[2]).merge(ev.rows[4].cells[2])
ab(ev); fixw(ev, [Cm(2.7), Cm(8.0), Cm(3.4)])
box.add_paragraph().paragraph_format.space_after = Pt(2)
bl("[순매수 경품 지급기준]", size=9, bold=True, after=2)
gt = box.add_table(rows=4, cols=3)
for j, hh in enumerate(["지급기준", "경품", "인원"]): ct(gt.rows[0].cells[j], hh, bold=True, shade="F2F2F2")
gdata = [("100만원 이상 ~ 500만원 미만", "투썸플레이스 아메리카노 1잔", "100명 한도"),
         ("500만원 이상 ~ 1천만원 미만", "모바일 문화상품권 1만원", "100명 한도"),
         ("1천만원 이상", "모바일 문화상품권 3만원", "100명 한도")]
for i, row in enumerate(gdata, 1):
    for j, v in enumerate(row): ct(gt.rows[i].cells[j], v)
ab(gt); fixw(gt, [Cm(5.2), Cm(5.8), Cm(3.1)])
box.add_paragraph().paragraph_format.space_after = Pt(2)
bl("* 동일 거래 상대방 제공 한도 : 50,000원", size=9, bold=True, after=1)

cbox = m.rows[3].cells[1]; cbox.text = ""
cst = cbox.add_table(rows=6, cols=5)
for j, hh in enumerate(["구분", "대상", "포상금(원)", "비용(원)", "비용부담"]): ct(cst.rows[0].cells[j], hh, bold=True, shade="F2F2F2")
ct(cst.rows[1].cells[0], "ISA 적립식"); ct(cst.rows[1].cells[1], "200명")
ct(cst.rows[1].cells[2], "모바일 문화상품권\n10,000"); ct(cst.rows[1].cells[3], "2,000,000", align=R)
ct(cst.rows[2].cells[0], "연금저축 적립식"); ct(cst.rows[2].cells[1], "200명")
ct(cst.rows[2].cells[2], "모바일 문화상품권\n10,000"); ct(cst.rows[2].cells[3], "2,000,000", align=R)
ct(cst.rows[3].cells[0], "ISA 순매수"); ct(cst.rows[3].cells[1], "300명")
ct(cst.rows[3].cells[2], "구간별 경품\n4,900~30,000"); ct(cst.rows[3].cells[3], "4,490,000", align=R)
ct(cst.rows[4].cells[0], "DC/IRP 순매수"); ct(cst.rows[4].cells[1], "300명")
ct(cst.rows[4].cells[2], "구간별 경품\n4,900~30,000"); ct(cst.rows[4].cells[3], "4,490,000", align=R)
ct(cst.rows[1].cells[4], "한국투자\n신탁운용")
cst.rows[1].cells[4].merge(cst.rows[2].cells[4]).merge(cst.rows[3].cells[4]).merge(cst.rows[4].cells[4])
ct(cst.rows[5].cells[0], "한국투자신탁운용 비용 총액", bold=True, shade="F2F2F2")
cst.rows[5].cells[0].merge(cst.rows[5].cells[1]).merge(cst.rows[5].cells[2])
ct(cst.rows[5].cells[3], "12,980,000", bold=True, align=R, shade="F2F2F2"); ct(cst.rows[5].cells[4], "", shade="F2F2F2")
ab(cst); fixw(cst, [Cm(2.9), Cm(2.0), Cm(3.4), Cm(2.8), Cm(1.9)])
p = cbox.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
sf(p.add_run("※ 투썸플레이스 아메리카노 1잔 단가 4,900원 기준"), 9, False)
p = cbox.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
sf(p.add_run("※ 상기 비용은 예상비용으로 실제 조건 충족 결과에 따라 달라질 수 있음"), 9, True)
ab(m); fixw(m, [Cm(1.6), Cm(15.8)])

para(doc, "", after=6)
para(doc, "3. 비용처리항목 : 판매부대비-프로모션", after=6)
for tx in ["※ 재산상 이익 제공 사전 신고용",
           "※ 동일 거래 상대방 제공 한도 : 50,000원",
           "※ 재산상 이익 제공 금액 (예상) : 12,980,000원",
           "※ 부당한 재산상 이익 제공 해당없음"]:
    para(doc, tx, size=10, after=2)
para(doc, "", after=4)
para(doc, "[첨부]", bold=True, after=2)
para(doc, "  1) 프로모션 참가 공문 1부.  끝.", after=2)

out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "iM증권", "한국투자신탁운용_시행문_프로모션 참가_iM증권(ACE 고배당주Plus커버드콜액티브 외 4종).docx")
os.makedirs(os.path.dirname(out), exist_ok=True)
doc.save(out); print("SAVED:", out)
