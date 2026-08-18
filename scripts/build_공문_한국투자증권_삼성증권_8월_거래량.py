# -*- coding: utf-8 -*-
"""한국투자증권·삼성증권 8월 거래량 이벤트 공문 - ACE 삼성전자SK하이닉스플러스채권혼합50 (0233A0)
기간 2026-08-25~09-07 (10영업일), 일 약정 6천만원 이상 일 추첨 50명, 1만원 모바일 문화상품권(인당 최대 5만원)
비용부담: 한국투자신탁운용 (증권사당 5,000,000원)
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "맑은 고딕"
L, C, R = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT


def sf(run, size=10, bold=False, color=None):
    run.font.size = Pt(size); run.bold = bold
    rpr = run._element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'): rf.set(qn(a), FONT)
    if color: run.font.color.rgb = color


def para(doc, text="", size=10, bold=False, align=None, after=3, color=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.2
    if text: sf(p.add_run(text), size, bold, color)
    return p


def bd(cell):
    tcPr = cell._tc.get_or_add_tcPr(); tcB = tcPr.find(qn('w:tcBorders'))
    if tcB is None: tcB = OxmlElement('w:tcBorders'); tcPr.append(tcB)
    for e in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000'); tcB.append(el)


def ct(cell, text, size=9, bold=False, align=C, shade=None):
    cell.text = ""
    for i, line in enumerate(str(text).split("\n")):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        sf(p.add_run(line), size, bold)
    cell.vertical_alignment = 1
    if shade:
        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), shade); cell._tc.get_or_add_tcPr().append(sh)
    bd(cell)


def ab(t):
    for row in t.rows:
        for c in row.cells: bd(c)


def fixw(t, widths):
    t.autofit = False; t.allow_autofit = False
    tblPr = t._tbl.tblPr; lay = tblPr.find(qn('w:tblLayout'))
    if lay is None: lay = OxmlElement('w:tblLayout'); tblPr.append(lay)
    lay.set(qn('w:type'), 'fixed')
    for row in t.rows:
        for c, w in zip(row.cells, widths): c.width = w
    for col, w in zip(t.columns, widths): col.width = w


def build(broker):
    doc = Document(); s = doc.sections[0]
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.5); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    para(doc, "[ 한국투자신탁운용 레터헤드 로고 자리 ]", size=9, align=C, color=RGBColor(0x99, 0x99, 0x99), after=6)
    t = doc.add_table(rows=2, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    ct(t.rows[0].cells[0], "우) 07320 서울특별시 영등포구 여의대로 24", align=L)
    ct(t.rows[1].cells[0], "전 화 : 02-2055-5348    / FAX :        / 문의 안내 : ETF마케팅부 채보라[부서원]", align=L)
    ab(t); para(doc, "", after=6)
    for k, v in [("문서번호", "한국투자신탁운용-2026-____   (신규 채번)"),
                 ("발신일자", "2026-08-__   (발송일 기입)"),
                 ("수    신", f"{broker} 대표이사"),
                 ("참    조", ""),
                 ("제    목", f"{broker} 거래고객 대상 ACE ETF 거래량 이벤트 제안의 건 (ACE 삼성전자SK하이닉스플러스채권혼합50)")]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.2
        sf(p.add_run(k + "        "), 10, True); sf(p.add_run(v), 10, False)
    p = doc.add_paragraph(); pPr = p._p.get_or_add_pPr()
    pb = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom'); bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '12'); bt.set(qn('w:space'), '1'); bt.set(qn('w:color'), '1F3864'); pb.append(bt); pPr.append(pb)
    para(doc, "", after=8)
    para(doc, "        귀사의 무궁한 발전을 기원드리며, 아래와 같이 당사 ETF 상품에 대한 판매 활성화 차원의 "
              "프로모션을 제안하오니 검토하여 주시기 바랍니다.", align=L, after=8)
    para(doc, "- 아    래 -", bold=True, align=C, after=8)
    para(doc, "1. 목적 : ACE ETF 상품에 대한 투자자의 인지도 제고 및 거래 활성화", after=6)
    para(doc, "2. 주요내용", after=4)

    m = doc.add_table(rows=4, cols=2); m.alignment = WD_TABLE_ALIGNMENT.CENTER
    ct(m.rows[0].cells[0], "구 분", size=10, bold=True, shade="F2F2F2"); ct(m.rows[0].cells[1], "세부 내용", size=10, bold=True, shade="F2F2F2")
    ct(m.rows[1].cells[0], "기 간", size=10, bold=True); ct(m.rows[1].cells[1], "2026년 8월 25일 ~ 2026년 9월 7일 / 10영업일", size=10, align=L)
    ct(m.rows[2].cells[0], "프로모션\n내용", size=10, bold=True)
    ct(m.rows[3].cells[0], "예상\n비용", size=10, bold=True)

    box = m.rows[2].cells[1]; box.text = ""

    def bl(text, size=10, bold=False, after=2, first=False):
        p = box.paragraphs[0] if first else box.add_paragraph()
        p.alignment = L; p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.15
        sf(p.add_run(text), size, bold); return p

    bl(f"1. 대상고객 : {broker} 거래고객", first=True)
    bl("2. 대상상품 : ACE 삼성전자SK하이닉스플러스채권혼합50", after=4)
    sp = box.add_table(rows=2, cols=7)
    heads = ["상품명", "위험등급", "총\n보수", "집합\n투자", "지정\n참가", "신탁", "일반\n사무"]
    for j, hh in enumerate(heads): ct(sp.rows[0].cells[j], hh, bold=True, shade="F2F2F2")
    row = ("ACE 삼성전자SK하이닉스\n플러스채권혼합50\n(0233A0)", "4등급\n(보통 위험)", "0.070", "0.049", "0.001", "0.010", "0.010")
    for j, v in enumerate(row): ct(sp.rows[1].cells[j], v, align=(L if j == 0 else C))
    ab(sp); fixw(sp, [Cm(4.4), Cm(2.4), Cm(1.6), Cm(1.6), Cm(1.6), Cm(1.6), Cm(1.6)])
    p = box.add_paragraph(); p.alignment = L; p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(2)
    sf(p.add_run("* 상장일 : 2026년 8월 25일 (신규상장)"), 9, False)
    bl("3. 주요내용 :", after=3)
    ev = box.add_table(rows=2, cols=3)
    for j, hh in enumerate(["구분", "주요내용", "기타"]): ct(ev.rows[0].cells[j], hh, bold=True, shade="F2F2F2")
    ct(ev.rows[1].cells[0], "거래량 이벤트", bold=True)
    ct(ev.rows[1].cells[1], "- 일 약정금액 6천만원 이상\n- 일 조건 충족 추첨 50명\n- 1만원 모바일 문화상품권 증정\n  (인당 최대 5만원)", align=L)
    ct(ev.rows[1].cells[2], "- 운용사 비용 부담")
    ab(ev); fixw(ev, [Cm(2.9), Cm(7.3), Cm(4.6)])
    box.add_paragraph().paragraph_format.space_after = Pt(2)
    bl("* 동일 거래 상대방 제공 한도 : 50,000원 (10,000원 × 최대 5회)", size=9, bold=True, after=1)

    cbox = m.rows[3].cells[1]; cbox.text = ""
    cst = cbox.add_table(rows=3, cols=6)
    for j, hh in enumerate(["구분", "대상", "영업일수", "포상금(원)", "비용(원)", "비용부담"]): ct(cst.rows[0].cells[j], hh, bold=True, shade="F2F2F2")
    ct(cst.rows[1].cells[0], "거래량 이벤트"); ct(cst.rows[1].cells[1], "일 추첨 50명"); ct(cst.rows[1].cells[2], "10")
    ct(cst.rows[1].cells[3], "모바일 문화상품권\n10,000"); ct(cst.rows[1].cells[4], "5,000,000", align=R); ct(cst.rows[1].cells[5], "한국투자\n신탁운용")
    ct(cst.rows[2].cells[0], "한국투자신탁운용 비용 총액", bold=True, shade="F2F2F2")
    cst.rows[2].cells[0].merge(cst.rows[2].cells[1]).merge(cst.rows[2].cells[2]).merge(cst.rows[2].cells[3])
    ct(cst.rows[2].cells[4], "5,000,000", bold=True, align=R, shade="F2F2F2"); ct(cst.rows[2].cells[5], "", shade="F2F2F2")
    ab(cst); fixw(cst, [Cm(3.1), Cm(2.4), Cm(1.7), Cm(3.3), Cm(2.4), Cm(1.9)])
    p = cbox.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
    sf(p.add_run("※ 상기 비용은 예상비용으로 실제 거래체결에 따른 결과 달라질 수 있음"), 9, True)
    ab(m); fixw(m, [Cm(1.6), Cm(15.8)])

    para(doc, "", after=12)
    para(doc, "한국투자신탁운용(주)", size=13, bold=True, align=C, after=1)
    para(doc, "대표이사 배 재 규   (직인)", size=12, bold=True, align=C, after=0)

    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out = os.path.join(root, broker,
                       f"한국투자신탁운용_공문_거래량 이벤트_{broker}_2026년 8월(ACE 삼성전자SK하이닉스플러스채권혼합50).docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out); print("SAVED:", out)


for broker in ("한국투자증권", "삼성증권"):
    build(broker)
