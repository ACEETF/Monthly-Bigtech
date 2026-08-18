# -*- coding: utf-8 -*-
"""
ACE 삼성전자SK하이닉스플러스채권혼합50 · 주요 투자 포인트(4~14p) HTML 빌더

- 템플릿(투자포인트.template.html)의 플레이스홀더에
  (1) 백업 데이터 xlsx에서 산출한 차트 데이터(JSON)
  (2) 기존 제안서 PPTX에서 추출한 승인 차트 이미지 / ACE ETF 로고 (base64)
  를 주입해 단일 self-contained HTML 을 생성한다.

사용법:
    python3 scripts/build_투자포인트_html.py \
        --xlsx  <백업데이터.xlsx> \
        --pptx  <ACE K반도체TOP2 제안서.pptx> \
        --out   제안서/ACE_삼성전자SK하이닉스플러스채권혼합50/투자포인트.html
"""

import argparse
import base64
import datetime
import json
import os

import openpyxl
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DEFAULT_TPL = os.path.join(
    ROOT, "제안서", "ACE_삼성전자SK하이닉스플러스채권혼합50", "투자포인트.template.html"
)


# --------------------------------------------------------------------------
# 1) 차트 데이터 (백업 데이터 xlsx)
# --------------------------------------------------------------------------
def build_chart_data(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out = {}

    # 02_Plus_Calc : 일자(A) / Plus 누적수익률(N=14) / TOP2 누적수익률(U=21), 데이터 시작 16행
    ws = wb["02_Plus_Calc"]
    monthly = {}
    for r in range(16, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not isinstance(d, datetime.datetime):
            continue
        plus, top2 = ws.cell(r, 14).value, ws.cell(r, 21).value
        if plus is None or top2 is None:
            continue
        monthly[(d.year, d.month)] = [
            "%d-%02d" % (d.year, d.month), round(plus * 100, 2), round(top2 * 100, 2)
        ]
    out["plus_monthly"] = [v for _, v in sorted(monthly.items())]

    # 04_PEER_Calc : 성과지표 블록 (AC~AG = 29~33열, 4~7행)
    ws = wb["04_PEER_Calc"]
    out["peer_names"] = [ws.cell(3, c).value for c in range(29, 34)]
    out["peer"] = {
        key: [ws.cell(row, c).value for c in range(29, 34)]
        for key, row in (("cum", 4), ("cagr", 5), ("vol", 6), ("mdd", 7))
    }
    return out


# --------------------------------------------------------------------------
# 2) 이미지 자산 (기존 제안서 PPTX)
# --------------------------------------------------------------------------
# 기존 ACE K반도체TOP2+ 제안서에서 재사용하는 승인 차트 이미지
#   slide 4  : 빅테크 분기 CAPEX(WSJ) / OpenRouter 토큰 사용량(MacroMicro)
#   slide 10 : 삼성전기 FC-BGA 매출 추이 / 전사 영업이익 추이(대신증권 추정)
#   slide 1  : ACE ETF 로고
ASSET_MAP = {
    "__IMG_TOKEN__": (4, 46670),
    "__IMG_CAPEX__": (4, 32051),
    "__IMG_FCBGA__": (10, 87609),
    "__IMG_OP__": (10, 105938),
    "__IMG_LOGO__": (1, 6011),
}


def build_assets(pptx_path):
    prs = Presentation(pptx_path)
    pool = {}
    for slide_no in sorted({s for s, _ in ASSET_MAP.values()}):
        for shape in prs.slides[slide_no - 1].shapes:
            if shape.shape_type == 13:  # PICTURE
                img = shape.image
                pool[(slide_no, len(img.blob))] = (img.ext, img.blob)

    assets = {}
    for token, key in ASSET_MAP.items():
        if key not in pool:
            raise SystemExit(
                "이미지 자산을 찾지 못했습니다: %s (slide=%d, size=%d)" % (token, key[0], key[1])
            )
        ext, blob = pool[key]
        mime = "image/jpeg" if ext in ("jpg", "jpeg") else "image/png"
        assets[token] = "data:%s;base64,%s" % (mime, base64.b64encode(blob).decode())
    return assets


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="백업 데이터 xlsx")
    ap.add_argument("--pptx", required=True, help="기존 ACE K반도체TOP2+ 제안서 pptx")
    ap.add_argument("--template", default=DEFAULT_TPL)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    with open(args.template, encoding="utf-8") as fp:
        html = fp.read()

    data = build_chart_data(args.xlsx)
    html = html.replace("__CHART_DATA__", json.dumps(data, ensure_ascii=False))

    for token, uri in build_assets(args.pptx).items():
        html = html.replace(token, uri)

    left = [t for t in list(ASSET_MAP) + ["__CHART_DATA__"] if t in html]
    if left:
        raise SystemExit("치환되지 않은 플레이스홀더: %s" % ", ".join(left))

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as fp:
        fp.write(html)

    print("생성 완료: %s (%.1f KB)" % (args.out, os.path.getsize(args.out) / 1024))
    print("  · 월별 성과 데이터 %d개월 / PEER %s" % (len(data["plus_monthly"]), ", ".join(data["peer_names"])))


if __name__ == "__main__":
    main()
